"""Markdown -> Word (.docx) exporter.

Ported from the validated _md2docx.py. Critically, all text is written through
python-docx `run.text`, which XML-escapes `< > &` automatically — this removes
the original bug where a literal `BMI<18.5` corrupted the document.xml.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

from ..core.image_fetcher import fetch_one

_HEAD_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_IMG_LINE_RE = re.compile(r"^\s*!\[([^\]]*)\]\(([^)]+)\)\s*$")
_QUOTE_RE = re.compile(r"^>\s?(.*)$")
_BULLET_RE = re.compile(r"^(\s*)[-*]\s+(.*)$")
_NUMBER_RE = re.compile(r"^(\s*)\d+\.\s+(.*)$")
_BOLD_RE = re.compile(r"(\*\*.+?\*\*)")


def _add_runs(paragraph, text: str) -> None:
    """Add text to a paragraph, turning **bold** segments into bold runs."""
    for seg in _BOLD_RE.split(text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
            run = paragraph.add_run(seg[2:-2])
            run.bold = True
        else:
            paragraph.add_run(seg)


def _add_image(doc: Document, alt: str, url: str, images_dir: Path, width_inch: float,
               md_dir: Path | None = None) -> None:
    if url.startswith("http://") or url.startswith("https://"):
        local = fetch_one(url, str(images_dir), name=f"img{len(list(images_dir.glob('*')))+1:03d}")
        if not local:
            # graceful: leave a caption-only paragraph
            p = doc.add_paragraph()
            _add_runs(p, f"[图片下载失败: {alt}]")
            return
        url = local
    else:
        # local reference: resolve relative to the markdown file's directory
        p = Path(url)
        if not p.is_absolute() and md_dir is not None:
            p = (md_dir / p).resolve()
        url = str(p)
    try:
        doc.add_picture(str(url), width=Inches(width_inch))
        if alt:
            cap = doc.add_paragraph()
            cap.alignment = 1  # center
            r = cap.add_run(alt)
            r.italic = True
            r.font.size = Pt(8)
    except Exception:
        p = doc.add_paragraph()
        _add_runs(p, f"[无法嵌入图片: {alt}]")


def md_to_docx(
    md_path: str,
    out_path: str,
    images_dir: str | None = None,
    width_inch: float = 6.0,
) -> str:
    """Convert a Markdown file to a Word document. Returns out_path."""
    text = Path(md_path).read_text(encoding="utf-8")
    out_path = str(out_path)
    images_dir = Path(images_dir) if images_dir else Path(out_path).with_suffix("") / "_images"

    doc = Document()
    for raw in text.splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("<!--") and stripped.endswith("-->"):
            continue  # skip split comments

        m = _HEAD_RE.match(line)
        if m:
            h = doc.add_heading(level=len(m.group(1)))
            _add_runs(h, m.group(2).strip())
            continue

        if _IMG_LINE_RE.match(stripped):
            alt, url = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped).groups()
            _add_image(doc, alt, url.strip(), images_dir, width_inch, md_dir=Path(md_path).parent)
            continue

        m = _QUOTE_RE.match(line)
        if m:
            p = doc.add_paragraph()
            p.style = "Intense Quote"
            _add_runs(p, m.group(1))
            continue

        m = _BULLET_RE.match(line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, m.group(2))
            continue

        m = _NUMBER_RE.match(line)
        if m:
            p = doc.add_paragraph(style="List Number")
            _add_runs(p, m.group(2))
            continue

        p = doc.add_paragraph()
        _add_runs(p, line)

    doc.save(out_path)
    return out_path


def result_to_docx(result, out_path: str, **kw) -> str:
    """Convenience: write a ConversionResult.markdown to a temp .md then to docx."""
    tmp_md = Path(out_path).with_suffix(".tmp.md")
    tmp_md.write_text(result.markdown, encoding="utf-8")
    try:
        return md_to_docx(str(tmp_md), out_path, **kw)
    finally:
        tmp_md.unlink(missing_ok=True)
