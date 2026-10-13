import os
import tempfile
import zipfile
from xml.dom import minidom

from PIL import Image

from docflow.exporters.docx import md_to_docx


def _write_md(path: str) -> None:
    png = os.path.join(os.path.dirname(path), "pixel.png")
    Image.new("RGB", (20, 20), (200, 30, 30)).save(png, "PNG")
    md = (
        "# 第一章 绪论\n"
        "\n"
        "这是**加粗**文本，下面是一个不等式 BMI<18.5 表示体重过低。\n"
        "\n"
        "> 这是引用块\n"
        "\n"
        "- 列表项一\n"
        "- 列表项二\n"
        "\n"
        "![示意图](pixel.png)\n"
    )
    with open(path, "w", encoding="utf-8") as f:
        f.write(md)


def test_md_to_docx_valid_and_escaped():
    with tempfile.TemporaryDirectory() as td:
        md = os.path.join(td, "in.md")
        _write_md(md)
        out = os.path.join(td, "out.docx")
        md_to_docx(md, out, images_dir=td)

        # 1) file is a valid docx (zip)
        assert zipfile.is_zipfile(out), "docx must be a valid zip"

        # 2) document.xml parses as XML (the old bug produced invalid XML)
        xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
        minidom.parseString(xml)

        # 3) the literal '<' in 'BMI<18.5' must be escaped to &lt;
        assert "BMI&lt;18.5" in xml, "unescaped '<' would corrupt the document"

        # 4) headings + image present
        from docx import Document

        d = Document(out)
        headings = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
        assert "第一章 绪论" in headings
        assert len(d.inline_shapes) == 1
        # bold run exists
        bold_found = any(r.bold for p in d.paragraphs for r in p.runs)
        assert bold_found
